# -*- coding: utf-8 -*-
import os
import re
import sys
import time
import traceback
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS

import pymupdf as fitz
import pytesseract
from PIL import Image
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

if os.name == "nt":
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

app = Flask(__name__)
CORS(app)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)
app.config["MAX_CONTENT_LENGTH"] = 4 * 1024 ** 3

sys.path.insert(0, BASE_DIR)
import smart_extract as se

COLUMNS = [
    "رقم الفاتورة", "تاريخ الإصدار", "الرقم المرجعي للعميل", "اسم العميل",
    "كود/باركود المنتج", "اسم المنتج", "الكمية", "سعر الوحدة",
    "شامل الضريبة؟", "الموقع", "الضريبة%", "نسبة الخصم", "قيمة الخصم"
]

processing_state = {
    "running": False, "total": 0, "done": 0, "errors": 0,
    "current_file": "", "results": [], "start_time": None,
    "elapsed": "0s", "output_file": None, "error_log": []
}


@app.route("/")
def index():
    return send_from_directory(BASE_DIR, "index.html")


UPLOAD_TOKEN = "UPLOADS"


def _scan_dir(folder):
    extensions = {'.pdf', '.png', '.jpg', '.jpeg', '.tiff', '.tif', '.bmp', '.webp'}
    files = []
    for root, dirs, fnames in os.walk(folder):
        for f in fnames:
            ext = os.path.splitext(f)[1].lower()
            if ext in extensions:
                fp = os.path.join(root, f)
                files.append({"path": fp, "name": f, "size": os.path.getsize(fp), "ext": ext})
    files.sort(key=lambda x: x["name"])
    return files


@app.route("/api/upload", methods=["POST"])
def upload_files():
    uploaded = request.files.getlist("files")
    if not uploaded:
        return jsonify({"error": "لا توجد ملفات مرفوعة"}), 400
    saved = []
    total = 0
    existing = set(os.listdir(UPLOAD_DIR))
    for f in uploaded:
        name = os.path.basename(f.filename or "")
        if not name:
            continue
        ext = os.path.splitext(name)[1].lower()
        if ext not in {'.pdf', '.png', '.jpg', '.jpeg', '.tiff', '.tif', '.bmp', '.webp'}:
            continue
        base, e = os.path.splitext(name)
        n, tries = name, 1
        while n in existing:
            n = "%s_%d%s" % (base, tries, e)
            tries += 1
        dest = os.path.join(UPLOAD_DIR, n)
        f.save(dest)
        existing.add(n)
        total += os.path.getsize(dest)
        saved.append(n)
    return jsonify({"count": len(saved), "names": saved, "total_size_mb": round(total / (1024 * 1024), 1)})


@app.route("/api/clear-uploads", methods=["POST"])
def clear_uploads():
    removed = 0
    for f in os.listdir(UPLOAD_DIR):
        try:
            os.remove(os.path.join(UPLOAD_DIR, f))
            removed += 1
        except Exception:
            pass
    return jsonify({"removed": removed})


@app.route("/api/scan-folder", methods=["POST"])
def scan_folder():
    data = request.get_json()
    folder = data.get("folder", "").strip()
    if folder == UPLOAD_TOKEN:
        files = _scan_dir(UPLOAD_DIR)
    else:
        if not folder or not os.path.isdir(folder):
            return jsonify({"error": "المجلد غير موجود: " + folder}), 400
        files = _scan_dir(folder)
    total_size = sum(f["size"] for f in files)
    return jsonify({"files": files, "count": len(files), "total_size_mb": round(total_size / (1024 * 1024), 1)})


@app.route("/api/process", methods=["POST"])
def start_processing():
    if processing_state["running"]:
        return jsonify({"error": "المعالجة قيد التشغيل بالفعل"}), 400

    data = request.get_json()
    folder = data.get("folder", "").strip()
    workers = data.get("workers", 4)
    lang = data.get("lang", "ara+eng")

    if folder == UPLOAD_TOKEN:
        files = [os.path.join(UPLOAD_DIR, f) for f in sorted(os.listdir(UPLOAD_DIR))
                 if os.path.splitext(f)[1].lower() in {'.pdf', '.png', '.jpg', '.jpeg', '.tiff', '.tif', '.bmp', '.webp'}]
    else:
        if not folder or not os.path.isdir(folder):
            return jsonify({"error": "المجلد غير موجود"}), 400
        extensions = {'.pdf', '.png', '.jpg', '.jpeg', '.tiff', '.tif', '.bmp', '.webp'}
        files = []
        for root, dirs, fnames in os.walk(folder):
            for f in fnames:
                ext = os.path.splitext(f)[1].lower()
                if ext in extensions:
                    files.append(os.path.join(root, f))
        files.sort()

    if not files:
        return jsonify({"error": "لا توجد ملفات"}), 400

    processing_state.update({
        "running": True, "total": len(files), "done": 0, "errors": 0,
        "current_file": "", "results": [], "start_time": time.time(),
        "elapsed": "0s", "output_file": None, "error_log": []
    })

    import threading
    threading.Thread(target=process_files_thread, args=(files, workers, lang), daemon=True).start()
    return jsonify({"message": f"بدأت معالجة {len(files)} ملف", "total": len(files)})


@app.route("/api/process-status", methods=["GET"])
def get_status():
    s = processing_state
    if s["start_time"] and s["running"]:
        elapsed = time.time() - s["start_time"]
        mins, secs = divmod(int(elapsed), 60)
        hours, mins = divmod(mins, 60)
        s["elapsed"] = f"{hours}h {mins}m {secs}s" if hours else (f"{mins}m {secs}s" if mins else f"{secs}s")
        if s["done"] > 0:
            rate = s["done"] / elapsed
            remaining = (s["total"] - s["done"]) / rate if rate > 0 else 0
            rm, rs = divmod(int(remaining), 60)
            rh, rm = divmod(rm, 60)
            s["remaining"] = f"{rh}h {rm}m" if rh else (f"{rm}m {rs}s" if rm else f"{rs}s")
            s["rate"] = f"{rate:.1f}"
        else:
            s["remaining"] = "حساب..."
            s["rate"] = "0"
    return jsonify({
        "running": s["running"], "total": s["total"], "done": s["done"],
        "errors": s["errors"],
        "current_file": os.path.basename(s["current_file"]) if s["current_file"] else "",
        "elapsed": s.get("elapsed", "0s"), "remaining": s.get("remaining", ""),
        "rate": s.get("rate", ""),
        "output_file": s["output_file"],
        "error_count": len(s["error_log"]),
        "percent": round(s["done"] / s["total"] * 100, 1) if s["total"] > 0 else 0
    })


@app.route("/api/stop", methods=["POST"])
def stop_processing():
    processing_state["running"] = False
    return jsonify({"message": "تم الإيقاف"})


@app.route("/api/download", methods=["GET"])
def download_file():
    fname = request.args.get("file", "") or processing_state.get("output_file", "")
    if not fname:
        return jsonify({"error": "لا يوجد ملف"}), 404
    fpath = os.path.join(OUTPUT_DIR, fname) if not os.path.isabs(fname) else fname
    if not os.path.exists(fpath):
        return jsonify({"error": "الملف غير موجود"}), 404
    return send_file(fpath, as_attachment=True, download_name=fname)


@app.route("/api/errors", methods=["GET"])
def get_errors():
    return jsonify({"errors": processing_state["error_log"][-100:]})


def process_files_thread(files, workers, lang):
    try:
        all_results = []
        batch_size = 50

        with ProcessPoolExecutor(max_workers=min(workers, os.cpu_count() or 4)) as executor:
            for i in range(0, len(files), batch_size):
                if not processing_state["running"]:
                    break
                batch = files[i:i + batch_size]
                futures = {executor.submit(extract_file_worker, f, lang): f for f in batch if processing_state["running"]}

                for future in as_completed(futures):
                    fpath = futures[future]
                    processing_state["current_file"] = fpath
                    try:
                        result = future.result(timeout=300)
                        if result and result.get("rows"):
                            all_results.extend(result["rows"])
                    except Exception as e:
                        processing_state["errors"] += 1
                        processing_state["error_log"].append({"file": os.path.basename(fpath), "error": str(e)})
                    processing_state["done"] += 1

        if all_results and processing_state["running"]:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_name = f"invoices_{timestamp}.xlsx"
            save_to_excel(all_results, os.path.join(OUTPUT_DIR, output_name))
            processing_state["output_file"] = output_name
        processing_state["results"] = all_results
    except Exception as e:
        processing_state["error_log"].append({"file": "SYSTEM", "error": str(e)})
    finally:
        processing_state["running"] = False
        processing_state["current_file"] = ""


def extract_file_worker(filepath, lang):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return {"rows": process_pdf(filepath, lang)}
    else:
        return {"rows": process_image(filepath, lang)}


def _ocr_region(doc, page_idx, ax, ay, lang="ara+eng", dpi=250):
    """OCR a small region of a PDF page around the customer-value anchor."""
    try:
        page = doc[page_idx]
        r = page.rect
        # the name value usually sits on the same band, slightly below the label
        bx = max(0, ax - 70)
        by = max(0, ay - 28)
        w = min(310, r.width - bx)
        h = min(42, r.height - by)
        if w <= 15 or h <= 15:
            return ""
        clip = fitz.Rect(bx, by, bx + w, by + h)
        pix = page.get_pixmap(dpi=dpi, clip=clip, alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        for psm in ("7", "6"):
            ocr = pytesseract.image_to_string(img, lang=lang, config="--psm %s" % psm)
            best = ""
            best_score = 0
            for line in ocr.splitlines():
                line = line.strip().replace("\u200e", "").replace("\u200f", "").replace("\u0640", "")
                for tok in ("Name", "الاسم", "العملية"):
                    line = line.replace(tok, "")
                words = re.findall(r"[\u0600-\u06FF]{2,}", line)
                joined = " ".join(words)
                score = len(words)
                if score > best_score:
                    best_score = score
                    best = joined
            if best_score >= 2:
                return best
        return ""
    except Exception:
        return ""


def _ocr_customer_fallback(doc, header, lang):
    """If the text-layer customer name is unusable (reversed glyphs, garbage or
    missing), OCR the area the label-driven extractor pointed at."""
    name = (header.get("customer_name") or "").strip()
    ar = len(re.findall(r"[\u0600-\u06FF]", name))
    ax, ay = header.get("customer_anchor") or (None, None)
    leak_label = ("الاسم", "العميل", "اﻟﻔﺎﺗﻮرة", "اﻟﻌﻤﻴﻞ", "Name", "Customer", "Details",
                  "Information", "Contract", "Invoice", "Discount", "Data",
                  "الفاتورة", "رقم الفاتورة", "بيانات", "تفاصيل", "اضافية", "Add",
                  "شعار", "Heder")
    low = name.lower()
    suspicious = (ax is not None and ay is not None and (
        low.endswith(":") or low.startswith(":")
        or any(tok in name for tok in leak_label)
        or any("|" == ch for ch in name)
        or ar == 0 and len(name) < 5
        or ar == 0 and not re.search(r"[A-Za-z]{6,}", name)
        or name.startswith(leak_label)
        or len(name) < 6))
    if not suspicious:
        return name
    page_idx = 0
    for pg in range(min(len(doc), 3)):
        r = doc[pg].rect
        if 0 <= ay <= r.height:
            page_idx = pg
            break
    ocr = _ocr_region(doc, page_idx, ax, ay, lang=lang)
    ocr_ar = len(re.findall(r"[\u0600-\u06FF]", ocr))
    if ocr_ar >= 4 and (
            ocr_ar > ar or name.startswith(leak_label) or len(name) < 6):
        return ocr
    return name


def process_pdf(filepath, lang):
    rows = []
    doc = fitz.open(filepath)
    all_text_blocks = []

    for page_idx in range(len(doc)):
        page = doc[page_idx]
        dict_data = page.get_text("dict")
        for block in dict_data["blocks"]:
            if "lines" in block:
                for line in block["lines"]:
                    for span in line["spans"]:
                        text = span["text"].strip()
                        if text:
                            all_text_blocks.append({
                                "text": text,
                                "x": span["bbox"][0],
                                "y": span["bbox"][1],
                                "x2": span["bbox"][2],
                                "y2": span["bbox"][3],
                                "size": span["size"],
                                "font": span["font"],
                                "page": page_idx
                            })

    full_text = "\n".join([b["text"] for b in all_text_blocks])

    # Layout-adaptive extraction (primary path)
    try:
        header, products = se.extract_all(all_text_blocks, full_text)
    except Exception:
        header, products = {}, []
    header.setdefault("tax_percent", "15")

    # OCR the customer-name area if the text layer is unusable
    header["customer_name"] = _ocr_customer_fallback(doc, header, lang)

    for p in products:
        p.update(header)
        p["discount_value"] = p.get("discount", 0)

    rows = [product_to_row(p) for p in products]

    # Legacy text/image fallbacks only when nothing was extracted
    if not rows:
        text_only = "\n".join([b["text"] for b in all_text_blocks])
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            text = page.get_text("text")
            if text.strip() and len(text.strip()) > 20:
                parsed = parse_text_fallback(text, filepath)
                if parsed:
                    for p in parsed:
                        p.update(header)
                    rows.extend([product_to_row(p) for p in parsed])
            images = page.get_images(full=True)
            for img_info in images:
                try:
                    xref = img_info[0]
                    base_image = doc.extract_image(xref)
                    img_bytes = base_image["image"]
                    ext_img = base_image["ext"]
                    tmp = os.path.join(os.path.expanduser("~"), "AppData", "Local", "Temp", f"ocr_{page_idx}_{xref}.{ext_img}")
                    with open(tmp, "wb") as f:
                        f.write(img_bytes)
                    pil_img = Image.open(tmp)
                    if pil_img.width >= 100 and pil_img.height >= 50:
                        ocr_text = pytesseract.image_to_string(pil_img, lang=lang, config="--psm 6")
                        if ocr_text.strip():
                            parsed = parse_text_fallback(ocr_text, filepath)
                            for p in parsed:
                                p.update(header)
                            rows.extend([product_to_row(p) for p in parsed])
                    os.unlink(tmp)
                except:
                    pass

    doc.close()
    return rows


def extract_header(full_text, blocks, filepath=""):
    header = {
        "invoice_number": "",
        "date": "",
        "customer_name": "",
        "customer_code": "",
        "location": "",
        "tax_percent": "15",
    }

    inv_patterns = [
        r'(?:Invoice\s*Number|رقم\s*(?:الفاتورة|الفا?tورة))\s*(?:INV91|INV[-\s]?\d+|\S+)',
        r'\b(INV[-\s]?\d+)\b',
        r'\b(SI[-\s]?\d+)\b',
    ]

    for b in blocks:
        t = b["text"]
        for ip in inv_patterns:
            m = re.search(ip, t, re.IGNORECASE)
            if m:
                if m.lastindex and m.lastindex >= 1:
                    header["invoice_number"] = m.group(1).strip()
                else:
                    val = m.group(0).strip()
                    code_m = re.search(r'(INV[-\s]?\d+|SI[-\s]?\d+)', val, re.IGNORECASE)
                    if code_m:
                        header["invoice_number"] = code_m.group(1).strip()
                        break

    if not header["invoice_number"]:
        m = re.search(r'\b(INV\d+|SI\d+)\b', full_text, re.IGNORECASE)
        if m:
            header["invoice_number"] = m.group(1).strip()

    issue_date_found = False
    for b in blocks:
        t = b["text"]
        if re.search(r'(?:Issue|اصدار|اﺻﺪار)', t, re.IGNORECASE):
            next_blocks = [nb for nb in blocks if abs(nb["y"] - b["y"]) < 5 and nb["x"] > b["x"]]
            for nb in next_blocks:
                dm = re.search(r'(\d{4}[-/]\d{1,2}[-/]\d{1,2})', nb["text"])
                if dm:
                    d = dm.group(1)
                    parts = re.split(r'[-/]', d)
                    header["date"] = f"{parts[2]}/{parts[1]}/{parts[0]}"
                    issue_date_found = True
                    break
            if issue_date_found:
                break

    if not issue_date_found:
        date_patterns = [
            r'(?:Issue\s*Date|تاريخ\s*(?:الإصدار|الإصدار|الافتتاح))\s*(\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{2,4})',
            r'(\d{4}[-/]\d{1,2}[-/]\d{1,2})',
            r'(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})',
        ]
        for dp in date_patterns:
            dm = re.search(dp, full_text, re.IGNORECASE)
            if dm:
                d = dm.group(1)
                parts = re.split(r'[-/]', d)
                if len(parts) == 3 and len(parts[0]) == 4:
                    header["date"] = f"{parts[2]}/{parts[1]}/{parts[0]}"
                else:
                    header["date"] = d
                break

    for b in blocks:
        t = b["text"]
        if re.search(r'اسم\s*(العميل|العمييل|العمياء)', t) or re.search(r'Customer\s*Name', t, re.IGNORECASE):
            continue
        if re.search(r'(?:Name|الاسم)\s*$', t):
            continue

    customer_name_candidates = []
    for b in blocks:
        t = b["text"].strip()
        y = b["y"]
        x = b["x"]
        if y > 130 and y < 170 and x > 300:
            clean = re.sub(r'\s+', ' ', t).strip()
            if clean and len(clean) > 2:
                customer_name_candidates.append({"text": clean, "y": y, "x": x})

    for c in customer_name_candidates:
        t = c["text"]
        if re.search(r'Customer|العميل|Seller|اﻟﺒﺎﺋﻊ|Name|الاسم|Address|اﻟﻌﻨﻮان|VAT|اﻟﴬﻳﺐ|Additional|اﺿﺎﻓﻴﺔ', t):
            continue
        if len(t) > 3:
            header["customer_name"] = t
            break

    if not header["customer_name"]:
        for b in blocks:
            t = b["text"].strip()
            y = b["y"]
            x = b["x"]
            if y > 130 and y < 170 and x > 300:
                if re.search(r'(شركة|مختبر|مستشفى|مؤسسة|عيادة|مركز|اتحاد|سلام|اتصالات|العالمية)', t):
                    clean = re.sub(r'\s+', ' ', t).strip()
                    if len(clean) > 3:
                        header["customer_name"] = clean
                        break

    if not header["customer_name"]:
        m = re.search(r'(?:Customer|Bill To|Ship To)\s*[:\s]*(.+?)(?:\n|$)', full_text)
        if m:
            val = m.group(1).strip()
            if not re.match(r'^CUS\d', val):
                header["customer_name"] = val[:60]

    code_match = re.search(r'\b(CUS[-\s]?\d+)\b', full_text, re.IGNORECASE)
    if code_match:
        header["customer_code"] = code_match.group(1).strip()

    current_name = header.get("customer_name", "")
    garbled = bool(re.search(r'[\uFB50-\uFDCF\uFDF0-\uFDFF\uFE70-\uFEFF]', current_name))
    pdffont_broken = any(ch in current_name for ch in 'ﻻﻵﻹﻷﴍ﴾ﴽ') if current_name else False

    if (not current_name or garbled or pdffont_broken):
        try:
            ocr_doc = fitz.open(filepath)
            if len(ocr_doc) > 0:
                page = ocr_doc[0]
                dpi = 200
                scale = dpi / 72.0
                pix = page.get_pixmap(dpi=dpi)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

                cust_anchor = None
                for c in customer_name_candidates:
                    t = c["text"]
                    if len(t) > 3 and not re.search(r'Customer|العميل|Seller|اﻟﺒﺎﺋﻊ|Name|الاسم|Address|اﻟﻌﻨﻮان|VAT|اﻟﴬﻳﺐ|Additional|اﺿﺎﻓﻴﺔ|CUS\d', t):
                        cust_anchor = c
                        break
                if cust_anchor:
                    cx, cy = cust_anchor["x"], cust_anchor["y"]
                    crop = img.crop((int((cx - 10) * scale), int((cy - 15) * scale),
                                     int((cx + 180) * scale), int((cy + 30) * scale)))
                else:
                    crop = img.crop((int(370 * scale), int(130 * scale), int(575 * scale), int(175 * scale)))
                ocr_text = pytesseract.image_to_string(crop, lang='ara+eng', config='--psm 7')
                ocr_text = re.sub(r'^[\d\s.\-–—:؛,،]+', '', ocr_text).strip()
                ocr_text = re.sub(r'\s+', ' ', ocr_text).strip()
                if ocr_text and len(ocr_text) > 2:
                    skip_words = ['Customer', 'العميل', 'Seller', 'البائع', 'Name', 'الاسم', 'Address', 'العنوان']
                    if not any(sw in ocr_text for sw in skip_words):
                        header["customer_name"] = ocr_text
            ocr_doc.close()
        except Exception as e:
            print(f"[OCR] Error: {e}", flush=True)

    loc_match = re.search(r'(?:الموقع|الفرع|Location|Branch)\s*[:\s]*(.+?)(?:\n|$)', full_text)
    if loc_match:
        header["location"] = loc_match.group(1).strip()[:50]

    tax_found = False
    for b in blocks:
        t = b["text"].strip()
        if "Total Taxable" in t or "اﻻﺟﻤﺎﻟﻲ" in t:
            continue
        tm = re.match(r'^(\d{1,2})%$', t)
        if tm:
            val = int(tm.group(1))
            if val in (5, 10, 15, 20):
                header["tax_percent"] = str(val)
                tax_found = True
                break
    if not tax_found:
        header["tax_percent"] = "15"

    return header


def extract_products_structured(blocks, header):
    products = []

    table_header_y = None
    for b in blocks:
        t = b["text"].strip()
        if "Qty" in t or "الكمية" in t:
            table_header_y = b["y"]
            break

    if table_header_y is None:
        return products

    summary_keywords = ["Invoice Summary", "ملخص", "Total", "اﻻﺟﻤﺎﻟﻲ", "Discount", "الخصم",
                        "SubTotal", "Ninety", "ريال", "BANK", "البنك", "Account", "الحساب",
                        "IBAN", "Invoice Number", "Issue Date", "Due Date", "Customer", "Seller",
                        "Halala", "Saudi Riyal", "Services Description", "Unit Price", "Qty",
                        "Total Before", "VAT", "SubTotal", "Additional", "Address", "Name",
                        "Invoice Summary", "اﻟﻔﺎﺗﻮرة", "ﻣﻠﺨﺺ"]

    item_numbers = []
    for b in blocks:
        t = b["text"].strip()
        y = b["y"]
        x = b["x"]
        is_summary = any(kw in t for kw in summary_keywords)
        if is_summary:
            continue
        if x > 500 and re.match(r'^\d{1,3}$', t):
            num = int(t)
            if 1 <= num <= 500:
                item_numbers.append({"num": num, "block": b})

    if not item_numbers:
        return products

    prev_item_y = table_header_y
    prev_page = item_numbers[0]["block"]["page"] if item_numbers else None

    for idx, item in enumerate(item_numbers):
        num = item["num"]
        item_y = item["block"]["y"]
        item_page = item["block"]["page"]

        if prev_page is not None and prev_page != item_page:
            prev_item_y = 0
        prev_page = item_page

        next_y = None
        if idx + 1 < len(item_numbers):
            next_y = item_numbers[idx + 1]["block"]["y"]
        else:
            next_y = item_y + 200

        desc_blocks = [b for b in blocks
                       if b["page"] == item["block"]["page"]
                       and prev_item_y < b["y"] < item_y - 2
                       and b["text"].strip()]

        midpoint = (prev_item_y + item_y) / 2
        own_desc_blocks = [b for b in desc_blocks if b["y"] > midpoint]

        num_blocks = [b for b in blocks
                      if b["page"] == item["block"]["page"]
                      and item_y - 3 <= b["y"] <= item_y + 3
                      and b["text"].strip()]

        below_mid = (item_y + next_y) / 2
        below_desc = [b for b in blocks
                      if b["page"] == item["block"]["page"]
                      and item_y + 3 < b["y"] < below_mid
                      and b["text"].strip()]

        desc_parts = []
        qty = 0
        unit_price = 0
        total_before_vat = 0
        vat_amount = 0
        total_with_vat = 0

        for b in own_desc_blocks:
            t = b["text"].strip()
            is_summary = any(kw in t for kw in summary_keywords)
            if not is_summary and len(t) > 3 and not re.match(r'^[\d\s,\.]+$', t):
                desc_parts.append(t)

        for b in num_blocks:
            t = b["text"].strip()
            x = b["x"]
            if not t:
                continue
            is_summary = any(kw in t for kw in summary_keywords)
            if is_summary:
                continue
            if x > 500 and re.match(r'^\d{1,3}$', t):
                continue
            cleaned = t.replace(",", "").replace(" ", "").strip()
            if re.match(r'^\d[\d]*\.?\d*$', cleaned):
                try:
                    val = float(cleaned)
                except:
                    continue
                if val <= 0:
                    continue
                if 275 < x < 315:
                    qty = int(val) if val == int(val) else 0
                elif 225 < x <= 275:
                    unit_price = val
                elif 145 < x <= 225:
                    total_before_vat = val
                elif 85 < x <= 145:
                    vat_amount = val
                elif 10 < x <= 85:
                    total_with_vat = val

        for b in below_desc:
            t = b["text"].strip()
            x = b["x"]
            if not t:
                continue
            is_summary = any(kw in t for kw in summary_keywords)
            if is_summary:
                continue
            if x > 500 and re.match(r'^\d{1,3}$', t):
                continue
            cleaned = t.replace(",", "").replace(" ", "").strip()
            if re.match(r'^\d[\d]*\.?\d*$', cleaned):
                try:
                    val = float(cleaned)
                except:
                    continue
                if val <= 0:
                    continue
                if 275 < x < 315:
                    qty2 = int(val) if val == int(val) else 0
                    if qty == 0 and qty2 > 0:
                        qty = qty2
                elif 225 < x <= 275:
                    if unit_price == 0 and val > 0:
                        unit_price = val
                elif 145 < x <= 225:
                    if total_before_vat == 0 and val > 0:
                        total_before_vat = val
                elif 85 < x <= 145:
                    if vat_amount == 0 and val > 0:
                        vat_amount = val
                elif 10 < x <= 85:
                    if total_with_vat == 0 and val > 0:
                        total_with_vat = val
            elif len(t) > 3 and not re.match(r'^[\d\s,\.]+$', t):
                desc_parts.append(t)

        description = " ".join(desc_parts).strip()

        code = ""
        gap_skip = 15
        own_code_blocks = [b for b in desc_blocks if b["y"] > prev_item_y + gap_skip]
        for b in sorted(own_code_blocks, key=lambda b: b["y"]):
            t = b["text"].strip()
            cm = re.match(r'^([A-Z0-9]{2,15})\s*[-–—]', t)
            if cm:
                code = cm.group(1)
                break
        if not code:
            code_m = re.match(r'^([A-Z0-9]{2,15})\s*[-–—]', description)
            if code_m:
                code = code_m.group(1)

        if description and (qty > 0 or unit_price > 0 or total_before_vat > 0):
            if unit_price == 0 and qty > 0 and total_before_vat > 0:
                unit_price = round(total_before_vat / qty, 2)
            if total_before_vat == 0 and qty > 0 and unit_price > 0:
                total_before_vat = round(qty * unit_price, 2)
            if total_with_vat == 0 and total_before_vat > 0 and vat_amount > 0:
                total_with_vat = round(total_before_vat + vat_amount, 2)

            products.append({
                "item_number": num,
                "product_code": code,
                "product_name": description[:100],
                "quantity": qty,
                "unit_price": unit_price,
                "total_before_vat": total_before_vat,
                "vat_amount": vat_amount,
                "total_with_vat": total_with_vat,
            })

        prev_item_y = item_y

    return products


def extract_summary(full_text):
    summary = {}
    m_disc = re.search(r'(?:Discount|الخصم)\s*[\s:]*(\d[\d,]*\.?\d*)', full_text)
    if m_disc:
        summary["discount"] = float(m_disc.group(1).replace(",", ""))
    m_tax = re.search(r'Total\s*VAT.*?(\d[\d,]*\.?\d*)', full_text, re.IGNORECASE)
    if m_tax:
        summary["total_vat"] = float(m_tax.group(1).replace(",", ""))
    m_total = re.search(r'Total\s*Amount.*?(\d[\d,]*\.?\d*)', full_text, re.IGNORECASE)
    if m_total:
        summary["total_with_vat"] = float(m_total.group(1).replace(",", ""))
    m_taxable = re.search(r'Total\s*Taxable.*?(\d[\d,]*\.?\d*)', full_text, re.IGNORECASE)
    if m_taxable:
        summary["total_taxable"] = float(m_taxable.group(1).replace(",", ""))

    summary["tax_percent"] = "15"
    return summary


def extract_products_from_text(full_text, header):
    products = []
    lines = full_text.split("\n")

    item_pattern = re.compile(
        r'^(\d{1,3})\s*$'
    )
    code_pattern = re.compile(
        r'^([A-Z0-9]{3,10})\s*[-–—]\s*(.+)'
    )

    current_item = None
    current_desc = []
    current_nums = []

    for line in lines:
        line = line.strip()
        if not line:
            continue

        m_item = item_pattern.match(line)
        if m_item:
            if current_item is not None and current_desc:
                desc = " ".join(current_desc)[:100]
                nums = [float(n.replace(",", "")) for n in current_nums if n]
                code_m = re.match(r'^([A-Z0-9]{3,10})\s*[-–—]', desc)
                products.append({
                    "item_number": current_item,
                    "product_code": code_m.group(1) if code_m else "",
                    "product_name": desc,
                    "quantity": int(nums[0]) if len(nums) > 0 else 0,
                    "unit_price": nums[1] if len(nums) > 1 else 0,
                    "total_before_vat": nums[2] if len(nums) > 2 else 0,
                    "vat_amount": nums[3] if len(nums) > 3 else 0,
                    "total_with_vat": nums[4] if len(nums) > 4 else 0,
                })
            current_item = int(m_item.group(1))
            current_desc = []
            current_nums = []
            continue

        if current_item is not None:
            cleaned = line.replace(",", "").strip()
            if re.match(r'^\d[\d]*\.?\d*$', cleaned):
                current_nums.append(cleaned)
            elif len(line) > 5 and not re.match(r'^[\d\s,\.]+$', line):
                current_desc.append(line)

    if current_item is not None and current_desc:
        desc = " ".join(current_desc)[:100]
        nums = [float(n.replace(",", "")) for n in current_nums if n]
        code_m = re.match(r'^([A-Z0-9]{3,10})\s*[-–—]', desc)
        products.append({
            "item_number": current_item,
            "product_code": code_m.group(1) if code_m else "",
            "product_name": desc,
            "quantity": int(nums[0]) if len(nums) > 0 else 0,
            "unit_price": nums[1] if len(nums) > 1 else 0,
            "total_before_vat": nums[2] if len(nums) > 2 else 0,
            "vat_amount": nums[3] if len(nums) > 3 else 0,
            "total_with_vat": nums[4] if len(nums) > 4 else 0,
        })

    return products


def parse_text_fallback(text, source_file=""):
    products = []
    lines = [l.strip() for l in text.split("\n") if l.strip()]

    item_pattern = re.compile(r'^(\d{1,3})$')
    current_item = None
    current_desc = []
    current_nums = []

    for line in lines:
        m_item = item_pattern.match(line)
        if m_item:
            if current_item is not None and current_desc:
                desc = " ".join(current_desc)[:100]
                nums = [float(n.replace(",", "")) for n in current_nums if n]
                code_m = re.match(r'^([A-Z0-9]{3,10})\s*[-–—]', desc)
                products.append({
                    "item_number": current_item,
                    "product_code": code_m.group(1) if code_m else "",
                    "product_name": desc,
                    "quantity": int(nums[0]) if len(nums) > 0 else 0,
                    "unit_price": nums[1] if len(nums) > 1 else 0,
                    "total_before_vat": nums[2] if len(nums) > 2 else 0,
                    "vat_amount": nums[3] if len(nums) > 3 else 0,
                    "total_with_vat": nums[4] if len(nums) > 4 else 0,
                })
            current_item = int(m_item.group(1))
            current_desc = []
            current_nums = []
            continue

        if current_item is not None:
            cleaned = line.replace(",", "").strip()
            if re.match(r'^\d[\d]*\.?\d*$', cleaned):
                current_nums.append(cleaned)
            elif len(line) > 5 and not re.match(r'^[\d\s,\.]+$', line):
                current_desc.append(line)

    if current_item is not None and current_desc:
        desc = " ".join(current_desc)[:100]
        nums = [float(n.replace(",", "")) for n in current_nums if n]
        code_m = re.match(r'^([A-Z0-9]{3,10})\s*[-–—]', desc)
        products.append({
            "item_number": current_item,
            "product_code": code_m.group(1) if code_m else "",
            "product_name": desc,
            "quantity": int(nums[0]) if len(nums) > 0 else 0,
            "unit_price": nums[1] if len(nums) > 1 else 0,
            "total_before_vat": nums[2] if len(nums) > 2 else 0,
            "vat_amount": nums[3] if len(nums) > 3 else 0,
            "total_with_vat": nums[4] if len(nums) > 4 else 0,
        })

    return products


def process_image(filepath, lang):
    try:
        pil_img = Image.open(filepath)
        if pil_img.mode != "RGB":
            pil_img = pil_img.convert("RGB")
        ocr_text = pytesseract.image_to_string(pil_img, lang=lang, config="--psm 6")
        if ocr_text.strip():
            return parse_text_fallback(ocr_text, filepath)
    except Exception as e:
        raise e
    return []


def product_to_row(p):
    return [
        p.get("invoice_number", ""),
        p.get("date", ""),
        p.get("customer_code", ""),
        p.get("customer_name", ""),
        p.get("product_code", ""),
        p.get("product_name", ""),
        p.get("quantity", 0),
        p.get("unit_price", 0),
        "لا",
        p.get("location", ""),
        p.get("tax_percent", "15"),
        0,
        p.get("discount_value", 0)
    ]


def save_to_excel(rows, filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "الفواتير"

    header_font = Font(bold=True, color="FFFFFF", size=11, name="Arial")
    header_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0")
    )
    data_align = Alignment(horizontal="right", vertical="center")
    even_fill = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")

    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    ws.auto_filter.ref = f"A1:M{len(rows) + 1}"
    ws.freeze_panes = "A2"

    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = data_align
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = even_fill

    col_widths = [16, 14, 20, 35, 15, 55, 10, 14, 16, 30, 10, 12, 14]
    for i, w in enumerate(col_widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w

    ws.sheet_properties.tabColor = "1B4F72"
    wb.save(filepath)
    wb.close()


if __name__ == "__main__":
    print("=" * 60)
    print("  أداة استخراج فواتير PDF و الصور إلى Excel")
    print("  http://localhost:5000")
    print("=" * 60)
    app.run(debug=False, host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), threaded=True)
